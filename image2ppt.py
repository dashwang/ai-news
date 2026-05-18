#!/usr/bin/env python3
"""
图片转PPT工具
用法: python image2ppt.py <图片路径> [图片路径2] ...
"""

import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from PIL import Image

def images_to_ppt(image_paths, output_path="output.pptx"):
    """将多张图片转换成PPT"""
    
    prs = Presentation()
    prs.slide_width = Inches(13.333)  # 16:9 宽度
    prs.slide_height = Inches(7.5)    # 16:9 高度
    
    for img_path in image_paths:
        if not os.path.exists(img_path):
            print(f"⚠️ 图片不存在: {img_path}")
            continue
            
        # 添加空白幻灯片
        blank_slide_layout = prs.slide_layouts[6]  # 空白布局
        slide = prs.slides.add_slide(blank_slide_layout)
        
        # 获取图片尺寸
        with Image.open(img_path) as img:
            img_width, img_height = img.size
            # 保持宽高比
        
        # 计算显示尺寸（最大宽度10英寸）
        max_width = Inches(10)
        max_height = Inches(6)
        
        width_inches = min(max_width, Inches(img_width / 96))  # 96 DPI
        height_inches = width_inches * (img_height / img_width)
        
        if height_inches > max_height:
            height_inches = max_height
            width_inches = height_inches * (img_width / img_height)
        
        # 居中放置
        left = (prs.slide_width - width_inches) / 2
        top = (prs.slide_height - height_inches) / 2
        
        # 添加图片
        slide.shapes.add_picture(img_path, left, top, width=width_inches)
        
        # 添加标题（文件名）
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(3), Inches(0.5))
        tf = title_box.text_frame
        tf.text = os.path.splitext(os.path.basename(img_path))[0]
        tf.paragraphs[0].font.size = Pt(18)
        
        print(f"✅ 添加: {os.path.basename(img_path)}")
    
    prs.save(output_path)
    print(f"\n🎉 PPT已保存: {output_path}")
    return output_path

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("用法: python image2ppt.py <图片1.jpg> [图片2.jpg] ...")
        sys.exit(1)
    
    images = sys.argv[1:]
    output = "images_presentation.pptx"
    
    if "-o" in images:
        idx = images.index("-o")
        output = images[idx + 1]
        images = images[:idx] + images[idx+2:]
    
    images_to_ppt(images, output)
